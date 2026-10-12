"""Export du rapport d'analyse en JSON, PDF et Word."""

from __future__ import annotations

import hashlib
import json
from dataclasses import asdict
from datetime import datetime, timezone
from html import escape
from io import BytesIO

from .detection.signals import DetectionResult
from .verification.factcheck import VERDICT_LABELS, FactCheckReport


def build_report(
    *,
    filename: str,
    data: bytes,
    detection: DetectionResult | None,
    factcheck: FactCheckReport | None,
) -> dict:
    report: dict = {
        "genere_le": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "fichier": filename,
        "sha256": hashlib.sha256(data).hexdigest(),
        "taille_octets": len(data),
    }
    if detection is not None:
        probability = detection.probability
        report["detection_ia"] = {
            "type": detection.media_type,
            "verdict": detection.verdict,
            "probabilite_ia": round(probability, 3) if probability is not None else None,
            "indices": [asdict(s) for s in detection.signals],
            "metadonnees": detection.metadata,
        }
    if factcheck is not None:
        report["verification"] = factcheck.as_dict()
    return report


def to_json(report: dict) -> bytes:
    return json.dumps(report, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def _lines(report: dict) -> list[tuple[str, str]]:
    """Contenu commun aux exports : liste de (style, texte)."""
    lines = [
        ("title", "Rapport d'analyse Black Face"),
        ("p", f"Fichier : {report['fichier']}"),
        ("p", f"Empreinte SHA-256 : {report['sha256']}"),
        ("p", f"Généré le : {report['genere_le']}"),
    ]
    detection = report.get("detection_ia")
    if detection:
        probability = detection["probabilite_ia"]
        lines += [
            ("h1", "Détection de contenu généré par IA"),
            ("p", f"Verdict : {detection['verdict']}"
                  + (f" ({probability:.0%} de probabilité IA)" if probability is not None else "")),
        ]
        for signal in detection["indices"]:
            value = signal["error"] or signal["detail"]
            if signal["probability"] is not None:
                value = f"{signal['probability']:.0%} — {value}"
            lines.append(("li", f"{signal['name']} : {value}"))
    verification = report.get("verification")
    if verification:
        lines += [("h1", "Vérification des faits"), ("p", verification["resume"])]
        if verification["erreur"]:
            lines.append(("p", verification["erreur"]))
        for claim in verification["affirmations"]:
            lines.append(("h2", claim["claim"]))
            label = VERDICT_LABELS.get(claim["verdict"], claim["verdict"])
            lines.append(("p", f"Verdict : {label} (confiance {claim['confidence']})"))
            lines.append(("p", claim["answer"]))
            lines.append(("p", claim["explanation"]))
            for warning in claim["warnings"]:
                lines.append(("p", f"Avertissement : {warning}"))
            for source in claim["sources"]:
                item = source["evidence"]
                lines.append(("li", f"[{item['id']}] {source['stance']} — {item['domain']} — {item['title']} — {item['url']}"))
    lines.append(("small", "Les résultats sont des estimations et ne constituent pas une certification."))
    return lines


def to_pdf(report: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    mapping = {"title": "Title", "h1": "Heading2", "h2": "Heading4", "p": "BodyText", "li": "BodyText", "small": "Italic"}
    story = []
    for style, text in _lines(report):
        prefix = "• " if style == "li" else ""
        story.append(Paragraph(prefix + escape(text), styles[mapping[style]]))
        if style in {"h1", "title"}:
            story.append(Spacer(1, 2 * mm))
    buffer = BytesIO()
    SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm).build(story)
    return buffer.getvalue()


def to_docx(report: dict) -> bytes:
    from docx import Document

    document = Document()
    for style, text in _lines(report):
        if style == "title":
            document.add_heading(text, level=0)
        elif style == "h1":
            document.add_heading(text, level=1)
        elif style == "h2":
            document.add_heading(text, level=2)
        elif style == "li":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
